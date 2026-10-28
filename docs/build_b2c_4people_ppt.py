#!/usr/bin/env python3
"""프로젝트 II, 4인 팀. 쉬운 순 B2C 주제 10개."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W, H = 13.333, 7.5
FONT = "Malgun Gothic"
TOTAL = 6

NAVY = RGBColor(0x1B, 0x2E, 0x3C)
INK = RGBColor(0x1C, 0x28, 0x30)
CREAM = RGBColor(0xF3, 0xEF, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xE0, 0x6A, 0x2C)
TEAL = RGBColor(0x1C, 0x6B, 0x64)
BLUE = RGBColor(0x2A, 0x4C, 0x78)
MUTED = RGBColor(0x5C, 0x68, 0x70)
LINE = RGBColor(0xDD, 0xD4, 0xC6)
SOFT = RGBColor(0xF7, 0xF5, 0xF1)

OUTS = [
    "/workspace/docs/B2C-4인-쉬운순-주제10.pptx",
    "/workspace/docs/b2c-4people-easy-topics.pptx",
]

# 구현이 쉬운 순. 4명이 안 겹치게 나뉘고, 도로명 주소가 들어갈 자리가 분명한 주제.
PROJECTS = [
    ("1", "카페 배달", "MyBatis", "쉬움", TEAL,
     "메뉴를 고르고, 배달 상태만 바꿉니다.",
     "A 회원·주소   B 메뉴   C 주문   D 관리자"),
    ("2", "원데이 클래스", "JPA", "쉬움", TEAL,
     "회차를 고르고, 정원만 넘지 않게 받습니다.",
     "A 회원   B 클래스   C 신청   D 관리자"),
    ("3", "맛집 예약", "MyBatis", "쉬움", TEAL,
     "식당, 시간, 인원만 받으면 됩니다.",
     "A 회원   B 식당   C 예약   D 관리자"),
    ("4", "펫 용품몰", "MyBatis", "보통", ORANGE,
     "담고 주문하면 재고가 줄어듭니다. 배송지가 있습니다.",
     "A 회원·배송지   B 상품   C 주문   D 관리자"),
    ("5", "온라인 서점", "JPA", "보통", ORANGE,
     "펫 용품몰과 같고, 리뷰만 더합니다.",
     "A 회원·배송지   B 도서   C 주문   D 관리자"),
    ("6", "꽃·식물 구독", "JPA", "보통", ORANGE,
     "한 번 주문이 아니라 매달 배송입니다.",
     "A 회원·배송지   B 플랜   C 회차   D 관리자"),
    ("7", "세탁 수거", "API", "보통", ORANGE,
     "수거 주소를 받고, 상태를 다음 단계로 옮깁니다.",
     "A 회원·주소   B 요금   C 신청   D 관리자"),
    ("8", "중고 장터", "JPA", "어려움", BLUE,
     "파는 회원과 사는 회원이 둘 다 있습니다.",
     "A 회원   B 판매글   C 거래   D 관리자"),
    ("9", "스터디카페", "MyBatis", "어려움", BLUE,
     "같은 좌석, 같은 시간은 한 명만 됩니다.",
     "A 회원   B 좌석   C 예약   D 관리자"),
    ("10", "영화 예매", "JPA", "어려움", BLUE,
     "영화, 상영시간, 좌석이 한꺼번에 연결됩니다.",
     "A 회원   B 영화   C 예매   D 관리자"),
]


def set_run_font(run, size, bold, color):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False
    r_pr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = r_pr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(r_pr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def add_text(slide, x, y, w, h, text, size=16, bold=False, color=INK, align="left", anchor="top"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set("anchor", {"top": "t", "middle": "ctr", "bottom": "b"}[anchor])
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold, color)
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def bg(slide):
    rect(slide, 0, 0, W, H, CREAM)


def footer(slide, page):
    add_text(slide, 0.48, 7.16, 9, 0.26, "프로젝트 II  ·  4명  ·  1번부터 고르면 됩니다", size=12, color=MUTED)
    add_text(slide, 11.15, 7.16, 1.7, 0.26, f"{page}  /  {TOTAL}", size=12, color=MUTED, align="right")


def kicker(slide, text):
    rect(slide, 0.48, 0.34, 0.28, 0.1, ORANGE, radius=0.5)
    add_text(slide, 0.88, 0.24, 10, 0.3, text, size=14, bold=True, color=ORANGE, anchor="middle")


def title(slide, text):
    add_text(slide, 0.48, 0.5, 12.3, 0.55, text, size=30, bold=True, color=NAVY)


def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    rect(s, 0, 0, 0.22, H, ORANGE)
    rect(s, 0.5, 0.38, 7.55, 6.74, NAVY, radius=0.08)
    add_text(s, 0.95, 0.7, 6.7, 0.32, "프로젝트 II   ·   SPRING   ·   4명", size=14, bold=True, color=RGBColor(0xE7, 0xA0, 0x72))
    add_text(s, 0.95, 1.35, 6.7, 0.7, "1번이 가장 쉽습니다", size=22, color=RGBColor(0xD5, 0xDE, 0xE4))
    add_text(s, 0.95, 2.15, 6.8, 1.7, "4명이 코딩하는\nB2C 주제 10개", size=36, bold=True, color=WHITE)
    add_text(s, 0.95, 4.35, 6.6, 1.5, "쉬운 순입니다.\n관리자, 도로명 주소, 모바일, 고도화는\n모든 주제에 들어갑니다.", size=16, color=RGBColor(0xC5, 0xD0, 0xD6))

    add_text(s, 8.4, 0.42, 4.5, 0.36, "쉬운 순", size=16, bold=True, color=NAVY)
    for i, p in enumerate(PROJECTS):
        num, name, _stack, level, color, _why, _roles = p
        y = 0.9 + i * 0.6
        rect(s, 8.4, y, 4.45, 0.52, WHITE, line=LINE, radius=0.14)
        add_text(s, 8.5, y, 0.58, 0.52, num, size=16, bold=True, color=color, align="center", anchor="middle")
        add_text(s, 9.12, y, 2.4, 0.52, name, size=15, bold=True, color=INK, anchor="middle")
        add_text(s, 11.4, y, 1.28, 0.52, level, size=13, bold=True, color=color, align="right", anchor="middle")


def slide_common(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, "과목 공통")
    title(s, "주제만 바뀌고, 4명의 일은 같습니다")

    blocks = [
        (TEAL, "설계 문서", "추가 요구분석\nDB 설계\nUse Case · Sequence\n기능정의서 · 화면연결도"),
        (ORANGE, "반드시 구현", "사용자 화면\n관리자 페이지\n도로명 주소 검색\n모바일 브라우저"),
        (BLUE, "마무리", "고도화 1가지\n테스트 130시간"),
    ]
    for i, (color, head, body) in enumerate(blocks):
        x = 0.48 + i * 4.2
        rect(s, x, 1.35, 4.0, 2.85, WHITE, line=LINE, radius=0.08)
        rect(s, x, 1.35, 4.0, 0.1, color)
        add_text(s, x + 0.28, 1.58, 3.45, 0.45, head, size=22, bold=True, color=color)
        add_text(s, x + 0.28, 2.15, 3.45, 1.8, body, size=16, color=INK)

    roles = [
        ("A", "회원", "가입, 로그인\n주소 검색"),
        ("B", "상품", "메뉴, 책, 식당\n클래스 등록"),
        ("C", "신청", "주문 또는 예약\n규칙 하나 검사"),
        ("D", "관리자", "목록, 상태 변경\n고도화 1개"),
    ]
    for i, (mark, name, body) in enumerate(roles):
        x = 0.48 + i * 3.18
        rect(s, x, 4.45, 3.02, 2.35, WHITE, line=LINE, radius=0.08)
        rect(s, x + 0.2, 4.68, 0.58, 0.58, NAVY, radius=0.2)
        add_text(s, x + 0.2, 4.68, 0.58, 0.58, mark, size=18, bold=True, color=WHITE, align="center", anchor="middle")
        add_text(s, x + 0.92, 4.72, 1.9, 0.5, name, size=20, bold=True, color=NAVY, anchor="middle")
        add_text(s, x + 0.22, 5.45, 2.6, 1.1, body, size=15, color=INK)
    footer(s, 2)


def rank_row(slide, y, project):
    num, name, stack, _level, color, why, roles = project
    rect(slide, 0.48, y, 12.36, 1.05, WHITE, line=LINE, radius=0.1)
    rect(slide, 0.48, y, 0.12, 1.05, color)
    add_text(slide, 0.7, y, 0.7, 1.05, num, size=24, bold=True, color=color, anchor="middle")
    add_text(slide, 1.45, y + 0.08, 6.6, 0.38, name, size=20, bold=True, color=NAVY, anchor="middle")
    add_text(slide, 1.45, y + 0.44, 8.3, 0.28, why, size=13, color=INK, anchor="middle")
    add_text(slide, 1.45, y + 0.7, 8.5, 0.28, roles, size=13, bold=True, color=MUTED, anchor="middle")
    rect(slide, 11.15, y + 0.3, 1.42, 0.42, color, radius=0.4)
    add_text(slide, 11.15, y + 0.3, 1.42, 0.42, stack, size=13, bold=True, color=WHITE, align="center", anchor="middle")


def slide_ranks(prs, page, heading, label, rows):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, label)
    title(s, heading)
    y = 1.32
    for project in rows:
        rank_row(s, y, project)
        y += 1.14
    footer(s, page)


def slide_pick(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, "4명 추천")
    title(s, "처음이면 이 세 개 중 하나입니다")
    picks = [
        ("1", "카페 배달", TEAL, [
            ("A", "회원, 로그인, 배달 주소"),
            ("B", "메뉴와 옵션"),
            ("C", "주문, 접수에서 배달완료"),
            ("D", "주문 목록, 오늘 매출"),
        ], "고도화  오늘 많이 팔린 메뉴"),
        ("2", "원데이 클래스", TEAL, [
            ("A", "회원, 로그인"),
            ("B", "클래스와 회차, 공방 주소"),
            ("C", "신청, 정원 초과 거절"),
            ("D", "회차 마감 현황"),
        ], "고도화  마감 임박 표시"),
        ("4", "펫 용품몰", ORANGE, [
            ("A", "회원, 배송지 주소"),
            ("B", "상품, 카테고리, 재고"),
            ("C", "장바구니와 주문"),
            ("D", "주문 목록, 매출"),
        ], "고도화  재고 부족 목록"),
    ]
    for i, (num, name, color, roles, extra) in enumerate(picks):
        x = 0.48 + i * 4.2
        rect(s, x, 1.35, 4.0, 5.45, WHITE, line=LINE, radius=0.08)
        rect(s, x, 1.35, 4.0, 0.1, color)
        add_text(s, x + 0.24, 1.58, 1.1, 0.55, num, size=28, bold=True, color=color)
        add_text(s, x + 1.2, 1.62, 2.55, 0.5, name, size=22, bold=True, color=NAVY, anchor="middle")
        yy = 2.35
        for mark, text in roles:
            rect(s, x + 0.24, yy, 0.42, 0.42, color, radius=0.2)
            add_text(s, x + 0.24, yy, 0.42, 0.42, mark, size=13, bold=True, color=WHITE, align="center", anchor="middle")
            add_text(s, x + 0.78, yy, 3.0, 0.42, text, size=14, color=INK, anchor="middle")
            yy += 0.58
        rect(s, x + 0.24, 5.55, 3.52, 0.95, SOFT, radius=0.1)
        add_text(s, x + 0.4, 5.55, 3.2, 0.95, extra, size=15, bold=True, color=NAVY, anchor="middle")
    footer(s, 5)


def slide_hard(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, "뒤쪽 번호")
    title(s, "8번부터는 4명이 같은 규칙을 같이 만집니다")
    add_text(s, 0.48, 1.18, 12, 0.36, "1번부터 7번은 맡은 화면이 거의 겹치지 않습니다.", size=16, color=MUTED)
    rows = [
        ("8", "중고 장터", "회원이 두 역할입니다.", "한 사람은 판매 글을 올리고, 다른 사람은 그 글을 삽니다. 거래 상태도 같이 바꿉니다.", BLUE),
        ("9", "스터디카페", "좌석과 시간이 묶입니다.", "예약할 때 이미 잡힌 자리인지 확인해야 합니다. 이 검사는 C가 만들고 D의 관리 화면과 연결됩니다.", BLUE),
        ("10", "영화 예매", "연결이 세 단계입니다.", "영화 아래에 상영이 있고, 상영 아래에 좌석이 있습니다. B와 C가 같은 테이블을 자주 고칩니다.", BLUE),
    ]
    y = 1.7
    for num, name, point, why, color in rows:
        rect(s, 0.48, y, 12.36, 1.62, WHITE, line=LINE, radius=0.08)
        rect(s, 0.48, y, 0.12, 1.62, color)
        add_text(s, 0.78, y, 0.9, 1.62, num, size=28, bold=True, color=color, anchor="middle")
        add_text(s, 1.85, y + 0.22, 5.5, 0.45, name, size=24, bold=True, color=NAVY)
        add_text(s, 7.4, y + 0.26, 5.1, 0.4, point, size=16, bold=True, color=color, align="right")
        add_text(s, 1.85, y + 0.82, 10.5, 0.58, why, size=15, color=INK)
        y += 1.76
    footer(s, 6)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "4명이 코딩하는 B2C 주제 10개"
    prs.core_properties.subject = "프로젝트 II Spring Framework"
    prs.core_properties.author = "팀 프로젝트 가이드"
    slide_cover(prs)
    slide_common(prs)
    slide_ranks(prs, 3, "1번부터 5번, 먼저 보세요", "쉬운 순  1–5", PROJECTS[:5])
    slide_ranks(prs, 4, "6번부터 10번, 그다음입니다", "쉬운 순  6–10", PROJECTS[5:])
    slide_pick(prs)
    slide_hard(prs)
    for path in OUTS:
        prs.save(path)
        print(path)


if __name__ == "__main__":
    build()
